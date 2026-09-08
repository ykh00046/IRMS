"""자재 LOT 이력 라우터 (/lot-history 화면의 데이터).

조회 전용·무인증(배합 분석과 같은 정책 — 사내 공용 단말). 역추적(/blend/material-lot-trace)
은 blend_routes 에 그대로 두고, 화면만 이쪽으로 모았다.

엔드포인트:
    GET /blend/lot-history/families     가족(레시피 계보) 선택지 + 자재 목록
    GET /blend/lot-history              타임라인 + 교체 사건 (family 또는 material 필수)
    GET /blend/lot-history/export       같은 내용 Excel 2시트
"""

import io
import sqlite3
from datetime import datetime
from typing import Any

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse

from ..db import get_db
from ..services import lot_history_service


def _history_or_400(connection, family, material, start_date, end_date) -> dict[str, Any]:
    try:
        return lot_history_service.lot_history(
            connection, family=family, material=material,
            start_date=start_date, end_date=end_date,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


def build_router() -> APIRouter:
    router = APIRouter()

    @router.get("/blend/lot-history/families")
    def lot_history_families(connection: sqlite3.Connection = Depends(get_db)) -> dict[str, Any]:
        return lot_history_service.list_families(connection)

    @router.get("/blend/lot-history")
    def lot_history(
        family: str = Query(default="", max_length=200),
        material: str = Query(default="", max_length=200),
        start_date: str = Query(default="", max_length=10),
        end_date: str = Query(default="", max_length=10),
        connection: sqlite3.Connection = Depends(get_db),
    ) -> dict[str, Any]:
        return _history_or_400(connection, family, material, start_date, end_date)

    @router.get("/blend/lot-history/export")
    def lot_history_export(
        family: str = Query(default="", max_length=200),
        material: str = Query(default="", max_length=200),
        start_date: str = Query(default="", max_length=10),
        end_date: str = Query(default="", max_length=10),
        connection: sqlite3.Connection = Depends(get_db),
    ) -> StreamingResponse:
        """화면에서 본 그대로 — 타임라인(구간) 시트 + 교체 이력 시트."""
        data = _history_or_400(connection, family, material, start_date, end_date)
        from openpyxl import Workbook
        from openpyxl.styles import Font

        head = Font(bold=True)
        wb = Workbook()
        ws = wb.active
        ws.title = "LOT 구간"
        scope = []
        if data["family"]:
            scope.append(f"레시피: {data['family']['label']}")
        if data["material"]:
            scope.append(f"자재: {data['material']}")
        period = f"{data['start_date'] or '처음'} ~ {data['end_date'] or '지금'}"
        ws.append([" · ".join(scope) + f" · 기간 {period}"])
        ws["A1"].font = Font(bold=True, size=12)
        ws.append([])
        ws.append(["레시피", "자재", "품목코드", "자재 LOT", "첫 사용일", "첫 제품 LOT",
                   "마지막 사용일", "마지막 제품 LOT", "배합 수"])
        for cell in ws[3]:
            cell.font = head
        for row in data["rows"]:
            for seg in row["segments"]:
                ws.append([
                    row["family_label"], row["material_name"], row["material_code"] or "",
                    seg["lot"] or "(미입력)", seg["first_date"], seg["first_product_lot"],
                    seg["last_date"], seg["last_product_lot"], seg["record_count"],
                ])
        for col, width in zip("ABCDEFGHI", (18, 18, 14, 18, 12, 18, 12, 18, 9)):
            ws.column_dimensions[col].width = width

        ws2 = wb.create_sheet("교체 이력")
        ws2.append(["교체일", "레시피", "제품 LOT", "자재", "이전 LOT", "새 LOT", "작업자"])
        for cell in ws2[1]:
            cell.font = head
        for c in data["changes"]:
            ws2.append([c["work_date"], c["family_label"], c["product_lot"], c["material_name"],
                        c["prev_lot"] or "", c["new_lot"] or "", c["worker"]])
        for col, width in zip("ABCDEFG", (12, 18, 18, 18, 18, 18, 12)):
            ws2.column_dimensions[col].width = width

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        stamp = datetime.now().strftime("%Y%m%d")
        filename = f"lot_history_{stamp}.xlsx"
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )

    return router
