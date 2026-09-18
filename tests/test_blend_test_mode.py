"""시험 배합(test blend) 백엔드 계약 테스트 — docs/test-blend-design.md §2~§7.

시험 배합은 정식과 **같은 통제**(저울·LOT·편차·서명·수기·증량·폐기)로 기록되면서
정식 생산 통계·알림은 오염하지 않아야 한다. 그 두 성질은 서로 반대 방향이라
한쪽만 지켜지기 쉬우므로, 여기서 양쪽을 동시에 못 박는다:

저장(§4)  recipe_id NULL 강제 · 총량 = 목표량 합 · 비율 서버 산출 · 목표량 0 거부 ·
          기준 레시피 허용 편차 · 반응기 검사 건너뜀 · LOT "T-" 접두(정식과 순번 분리)
조회(§5)  test=all|only|exclude 필터와 건수 · viscosity_state 는 절대 missing 아님 ·
          recent-product-lots 기본 정식만 · 전체 Excel 백업 '시험' 열
격리(§6)  대시보드·배합 분석·LOT 이력 제외 / 자재 불출량 공개 API 는 **포함**
출력(§5)  배합일지 Excel 에 '시험' 표식 없음(외부 제출용 — 표식 금지)
페이지(§7) GET /blend/test 는 배합 화면을 시험 모드로 렌더
"""

from __future__ import annotations

import importlib
import io
import uuid
from datetime import date


def _client():
    import src.config as cfg
    import src.main as mainmod

    importlib.reload(cfg)
    importlib.reload(mainmod)
    from fastapi.testclient import TestClient

    return TestClient(mainmod.app)


def _uid() -> str:
    return uuid.uuid4().hex[:6].upper()


def _mgmt_client():
    """책임자 로그인 + CSRF 헤더 (tests/test_blend_save_integrity.py 와 동일 패턴)."""
    client = _client()
    assert client.post(
        "/api/auth/management-login", json={"username": "admin", "password": "admin"}
    ).status_code == 200

    def csrf():
        tok = client.cookies.get("csrftoken")
        return {"x-csrftoken": tok} if tok else {}

    return client, csrf


def _worker_session(client, csrf, worker: str) -> str:
    client.get("/api/blend/records")  # csrf 쿠키 확보
    client.post("/api/workers", json={"name": worker}, headers=csrf())
    client.post("/api/blend/session/login", json={"worker": worker}, headers=csrf())
    return worker


def _import_recipe(
    client, csrf, product, materials, *, tolerance_g=None, use_reactor=None
):
    header = "반제품명\t" + "\t".join(m[0] for m in materials)
    row = product + "\t" + "\t".join(str(m[1]) for m in materials)
    body: dict = {"raw_text": f"{header}\n{row}", "force": True}
    if tolerance_g is not None:
        body["tolerance_g"] = tolerance_g
    if use_reactor is not None:
        body["use_reactor"] = use_reactor
    res = client.post("/api/recipes/import", json=body, headers=csrf())
    assert res.status_code == 200, res.text
    return res.json()["created_ids"][0]


_WORK_DATE = "2026-09-18"


def _test_payload(product, *, rows, base_recipe_id=None, work_date=_WORK_DATE, **extra):
    """시험 저장 본문 — rows: [(자재명, 목표량, 실제량, LOT)] 또는 dict 행."""
    details = []
    for row in rows:
        if isinstance(row, dict):
            details.append(row)
            continue
        name, theory, actual, lot = row
        details.append({
            "material_name": name,
            "theory_amount": theory,
            "actual_amount": actual,
            "material_lot": lot,
        })
    body = {
        "is_test": True,
        "product_name": product,
        "worker": "무시됨",           # 서버는 작업자 세션 이름을 쓴다
        "work_date": work_date,
        "total_amount": 1,            # 무시됨 — 서버가 목표량 합으로 산출
        "details": details,
    }
    if base_recipe_id is not None:
        body["base_recipe_id"] = base_recipe_id
    body.update(extra)
    return body


def _save_test(client, csrf, product, **kwargs):
    return client.post(
        "/api/blend/records", json=_test_payload(product, **kwargs), headers=csrf()
    )


# ── §4 저장 ────────────────────────────────────────────────────────
def test_test_save_forces_null_recipe_and_derives_total_and_ratio():
    """시험 저장: recipe_id NULL 강제 · 총량 = 목표량 합 · 비율 = 목표량/합 · LOT "T-" 접두.

    같은 날 같은 이름의 정식 기록이 있어도 순번이 섞이지 않는다(접두가 달라 base 가 분리).
    """
    client, csrf = _mgmt_client()
    product = "가상제품" + _uid()
    _worker_session(client, csrf, "가상작업" + _uid())
    rid = _import_recipe(client, csrf, product, [("원료A", 60), ("원료B", 40)])

    # 같은 날 정식 기록 1건 — 순번 분리의 대조군.
    prod_res = client.post("/api/blend/records", json={
        "recipe_id": rid, "product_name": product, "worker": "무시됨",
        "work_date": _WORK_DATE, "total_amount": 100,
        "details": [
            {"material_name": "원료A", "actual_amount": 60, "material_lot": "LA"},
            {"material_name": "원료B", "actual_amount": 40, "material_lot": "LB"},
        ],
    }, headers=csrf())
    assert prod_res.status_code == 200, prod_res.text
    prod_lot = prod_res.json()["product_lot"]

    res = _save_test(
        client, csrf, product,
        rows=[("원료A", 30, 30, "TLA"), ("원료B", 70, 70, "TLB")],
        base_recipe_id=rid,
        recipe_id=rid,   # 보내와도 서버가 무시(NULL 강제)
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["is_test"] is True
    assert body["recipe_id"] is None
    assert body["base_recipe_id"] == rid
    assert body["total_amount"] == 100.0          # 1 이 아니라 목표량 합
    assert body["product_lot"].startswith("T-" + product)
    # 정식 순번과 분리 — 둘 다 같은 날 첫 번째(01)로 끝난다.
    assert prod_lot.endswith("01") and body["product_lot"].endswith("01")
    assert body["product_lot"] != prod_lot
    ratios = [d["ratio"] for d in body["details"]]
    assert ratios == [30.0, 70.0]
    theories = [d["theory_amount"] for d in body["details"]]
    assert theories == [30.0, 70.0]
    # 상세 조회에 기준 레시피 이름까지 실린다(§5).
    detail = client.get(f"/api/blend/records/{body['id']}")
    assert detail.status_code == 200, detail.text
    assert detail.json()["base_recipe_name"] == product


def test_test_save_rejects_nonpositive_theory():
    """목표량이 0 이하인 행은 400 — 시험의 목표량은 총량·비율의 유일한 근거다."""
    client, csrf = _mgmt_client()
    product = "목표량" + _uid()
    _worker_session(client, csrf, "목표량작업" + _uid())

    res = _save_test(
        client, csrf, product,
        rows=[("원료A", 30, 30, "TLA"), ("원료B", 0, 0, "TLB")],
    )
    assert res.status_code == 400, res.text
    assert "원료B" in res.text

    # 목표량 자체를 보내지 않은 경우도 같다.
    res2 = _save_test(client, csrf, product, rows=[
        {"material_name": "원료C", "actual_amount": 10, "material_lot": "TLC"},
    ])
    assert res2.status_code == 400, res2.text
    assert "원료C" in res2.text


def test_test_save_keeps_actual_and_lot_requirements():
    """§4.5 유지 — 실제량·자재 LOT 은 시험에서도 필수."""
    client, csrf = _mgmt_client()
    product = "통제유지" + _uid()
    _worker_session(client, csrf, "통제작업" + _uid())

    no_actual = _save_test(client, csrf, product, rows=[
        {"material_name": "원료A", "theory_amount": 10, "material_lot": "TLA"},
    ])
    assert no_actual.status_code == 400, no_actual.text
    assert "원료A" in no_actual.text

    no_lot = _save_test(client, csrf, product, rows=[
        {"material_name": "원료A", "theory_amount": 10, "actual_amount": 10},
    ])
    assert no_lot.status_code == 400, no_lot.text
    assert "LOT" in no_lot.text


def test_test_save_uses_base_recipe_tolerance():
    """허용 편차는 불러온 레시피(base_recipe_id)의 tolerance_g — 없으면 기본 0.05g."""
    client, csrf = _mgmt_client()
    product = "편차확인" + _uid()
    _worker_session(client, csrf, "편차작업" + _uid())
    rid = _import_recipe(
        client, csrf, product, [("원료A", 60), ("원료B", 40)], tolerance_g=0.5
    )

    # 0.3g 편차 — 기본(0.05g)이면 막히고, 기준 레시피(0.5g)면 통과해야 한다.
    ok = _save_test(
        client, csrf, product,
        rows=[("원료A", 30, 30.3, "TLA"), ("원료B", 70, 70, "TLB")],
        base_recipe_id=rid,
    )
    assert ok.status_code == 200, ok.text

    ng = _save_test(
        client, csrf, product,
        rows=[("원료A", 30, 30.3, "TLA"), ("원료B", 70, 70, "TLB")],
    )
    assert ng.status_code == 400, ng.text
    assert "0.05" in ng.text


def test_test_save_skips_reactor_requirement():
    """반응기 진행 제품을 시험해도 반응기를 묻지 않는다(§4.4) — reactor 는 NULL 로 저장."""
    client, csrf = _mgmt_client()
    product = "반응기확인" + _uid()
    _worker_session(client, csrf, "반응기작업" + _uid())
    rid = _import_recipe(
        client, csrf, product, [("원료A", 60), ("원료B", 40)], use_reactor=True
    )
    # 대조군: 정식 저장은 반응기 없이 400.
    prod = client.post("/api/blend/records", json={
        "recipe_id": rid, "product_name": product, "worker": "무시됨",
        "work_date": _WORK_DATE, "total_amount": 100,
        "details": [
            {"material_name": "원료A", "actual_amount": 60, "material_lot": "LA"},
            {"material_name": "원료B", "actual_amount": 40, "material_lot": "LB"},
        ],
    }, headers=csrf())
    assert prod.status_code == 400, prod.text
    assert "반응기" in prod.text

    res = _save_test(
        client, csrf, product,
        rows=[("원료A", 60, 60, "TLA"), ("원료B", 40, 40, "TLB")],
        base_recipe_id=rid,
        reactor=2,   # 보내와도 시험은 반응기를 저장하지 않는다
    )
    assert res.status_code == 200, res.text
    assert res.json()["reactor"] is None


def test_test_save_rejects_unknown_base_recipe():
    """base_recipe_id 는 recipes 에 실존해야 한다(§4.1) — 없으면 400."""
    client, csrf = _mgmt_client()
    product = "기준없음" + _uid()
    _worker_session(client, csrf, "기준작업" + _uid())
    res = _save_test(
        client, csrf, product,
        rows=[("원료A", 10, 10, "TLA")],
        base_recipe_id=99_999_999,
    )
    assert res.status_code == 400, res.text
    assert "기준 레시피" in res.text


def test_test_save_allows_unknown_material_and_rejects_wrong_code():
    """마스터에 없는 새 원재료는 이름만으로 통과(NULL 저장), 틀린 품목코드는 400(§4.2)."""
    client, csrf = _mgmt_client()
    product = "새원재료" + _uid()
    _worker_session(client, csrf, "새원료작업" + _uid())

    ok = _save_test(client, csrf, product, rows=[
        {"material_name": "미등록원료" + _uid(), "theory_amount": 12.5,
         "actual_amount": 12.5, "material_lot": "TX1"},
        # 같은 이름이 여러 행이어도 된다(분할 계량) — 레시피 대조가 없다.
        {"material_name": "분할원료", "theory_amount": 10, "actual_amount": 10,
         "material_lot": "TX2"},
        {"material_name": "분할원료", "theory_amount": 5, "actual_amount": 5,
         "material_lot": "TX2"},
    ])
    assert ok.status_code == 200, ok.text
    assert ok.json()["total_amount"] == 27.5
    assert [d["material_id"] for d in ok.json()["details"]] == [None, None, None]

    ng = _save_test(client, csrf, product, rows=[
        {"material_name": "코드위조원료", "material_code": "ZZ-NO-SUCH-CODE",
         "theory_amount": 10, "actual_amount": 10, "material_lot": "TX3"},
    ])
    assert ng.status_code == 400, ng.text
    assert "품목코드" in ng.text


def test_test_save_audit_records_is_test():
    """감사 로그의 blend_record_create details 에 시험 여부가 남는다(§4.5)."""
    client, csrf = _mgmt_client()
    product = "감사확인" + _uid()
    _worker_session(client, csrf, "감사작업" + _uid())
    res = _save_test(client, csrf, product, rows=[("원료A", 10, 10, "TLA")])
    assert res.status_code == 200, res.text
    lot = res.json()["product_lot"]

    logs = client.get("/api/admin/audit-logs", params={"limit": 200})
    assert logs.status_code == 200, logs.text
    entry = next(
        (
            row for row in logs.json()["items"]
            if row.get("action") == "blend_record_create"
            and row.get("target_label") == lot
        ),
        None,
    )
    assert entry is not None, logs.text
    details = entry.get("details")
    if isinstance(details, str):
        import json as _json
        details = _json.loads(details)
    assert details.get("is_test") is True, details


# ── §3·§5 LOT 미리보기 ────────────────────────────────────────────
def test_next_lot_preview_uses_test_prefix():
    """next-lot?test=1 은 "T-" 접두 규칙으로 미리 본다(§3)."""
    client, csrf = _mgmt_client()
    product = "미리보기" + _uid()
    res = client.get("/api/blend/next-lot", params={"product": product, "date": _WORK_DATE})
    assert res.status_code == 200, res.text
    assert res.json()["next_lot"].startswith(product)
    tres = client.get(
        "/api/blend/next-lot", params={"product": product, "date": _WORK_DATE, "test": 1}
    )
    assert tres.status_code == 200, tres.text
    assert tres.json()["next_lot"].startswith("T-" + product)


# ── §5 조회 필터·건수 ─────────────────────────────────────────────
def test_records_test_filter_and_counts():
    """test=all(기본)/only/exclude 와 total_available 이 같은 필터를 쓴다. 잘못된 값은 422."""
    client, csrf = _mgmt_client()
    product = "필터확인" + _uid()
    worker = "필터작업" + _uid()
    _worker_session(client, csrf, worker)
    rid = _import_recipe(client, csrf, product, [("원료A", 60), ("원료B", 40)])
    prod = client.post("/api/blend/records", json={
        "recipe_id": rid, "product_name": product, "worker": "무시됨",
        "work_date": _WORK_DATE, "total_amount": 100,
        "details": [
            {"material_name": "원료A", "actual_amount": 60, "material_lot": "LA"},
            {"material_name": "원료B", "actual_amount": 40, "material_lot": "LB"},
        ],
    }, headers=csrf())
    assert prod.status_code == 200, prod.text
    test_res = _save_test(
        client, csrf, product,
        rows=[("원료A", 30, 30, "TLA"), ("원료B", 70, 70, "TLB")],
        base_recipe_id=rid,
    )
    assert test_res.status_code == 200, test_res.text

    def query(**params):
        res = client.get("/api/blend/records", params={"worker": worker, **params})
        assert res.status_code == 200, res.text
        return res.json()

    all_rows = query()
    assert len(all_rows["items"]) == 2
    assert all_rows["total_available"] == 2
    # 기본 응답에 is_test·base_recipe_id 가 실린다.
    assert {r["is_test"] for r in all_rows["items"]} == {True, False}

    only = query(test="only")
    assert [r["is_test"] for r in only["items"]] == [True]
    assert only["total_available"] == 1
    assert only["items"][0]["base_recipe_id"] == rid
    assert only["items"][0]["product_lot"].startswith("T-")

    exclude = query(test="exclude")
    assert [r["is_test"] for r in exclude["items"]] == [False]
    assert exclude["total_available"] == 1

    bad = client.get("/api/blend/records", params={"test": "nope"})
    assert bad.status_code == 422, bad.text


def test_test_record_viscosity_state_never_missing():
    """시험 기록은 '점도 미입력(missing)' 이 되지 않는다(§5) — 정식은 그대로 missing."""
    client, csrf = _mgmt_client()
    product = "점도확인" + _uid()
    worker = "점도작업" + _uid()
    _worker_session(client, csrf, worker)
    rid = _import_recipe(client, csrf, product, [("원료A", 60), ("원료B", 40)])
    # 점도 관리 대상으로 등록 — 이래야 정식 기록이 missing 이 된다.
    vp = client.post(
        "/api/viscosity/products", json={"code": product, "name": product}, headers=csrf()
    )
    assert vp.status_code == 200, vp.text

    prod = client.post("/api/blend/records", json={
        "recipe_id": rid, "product_name": product, "worker": "무시됨",
        "work_date": _WORK_DATE, "total_amount": 100,
        "details": [
            {"material_name": "원료A", "actual_amount": 60, "material_lot": "LA"},
            {"material_name": "원료B", "actual_amount": 40, "material_lot": "LB"},
        ],
    }, headers=csrf())
    assert prod.status_code == 200, prod.text
    test_res = _save_test(
        client, csrf, product,
        rows=[("원료A", 30, 30, "TLA"), ("원료B", 70, 70, "TLB")],
        base_recipe_id=rid,
    )
    assert test_res.status_code == 200, test_res.text

    rows = client.get("/api/blend/records", params={"worker": worker}).json()["items"]
    states = {r["is_test"]: r["viscosity_state"] for r in rows}
    assert states[False] == "missing"
    assert states[True] == "na"


def test_recent_product_lots_separates_test_lots():
    """기본은 정식 LOT 만, test=1 이면 시험 LOT 만(§5)."""
    client, csrf = _mgmt_client()
    product = "최근LOT" + _uid()
    _worker_session(client, csrf, "최근작업" + _uid())
    rid = _import_recipe(client, csrf, product, [("원료A", 60), ("원료B", 40)])
    prod = client.post("/api/blend/records", json={
        "recipe_id": rid, "product_name": product, "worker": "무시됨",
        "work_date": _WORK_DATE, "total_amount": 100,
        "details": [
            {"material_name": "원료A", "actual_amount": 60, "material_lot": "LA"},
            {"material_name": "원료B", "actual_amount": 40, "material_lot": "LB"},
        ],
    }, headers=csrf())
    assert prod.status_code == 200, prod.text
    prod_lot = prod.json()["product_lot"]
    test_res = _save_test(
        client, csrf, product,
        rows=[("원료A", 30, 30, "TLA"), ("원료B", 70, 70, "TLB")],
        base_recipe_id=rid,
    )
    assert test_res.status_code == 200, test_res.text
    test_lot = test_res.json()["product_lot"]

    default = client.get("/api/blend/recent-product-lots", params={"names": product})
    assert default.status_code == 200, default.text
    lots = [it["lot"] for it in default.json()["items"].get(product, [])]
    assert lots == [prod_lot]

    only_test = client.get(
        "/api/blend/recent-product-lots", params={"names": product, "test": 1}
    )
    assert only_test.status_code == 200, only_test.text
    tlots = [it["lot"] for it in only_test.json()["items"].get(product, [])]
    assert tlots == [test_lot]


def test_export_all_has_test_column():
    """전체 Excel 백업은 시험 기록을 포함하되 '시험' 열로 구분한다(§5)."""
    from openpyxl import load_workbook

    client, csrf = _mgmt_client()
    product = "백업확인" + _uid()
    worker = "백업작업" + _uid()
    _worker_session(client, csrf, worker)
    res = _save_test(client, csrf, product, rows=[("원료A", 10, 10, "TLA")])
    assert res.status_code == 200, res.text

    export = client.get("/api/blend/records/export-all", params={"worker": worker})
    assert export.status_code == 200, export.text
    ws = load_workbook(io.BytesIO(export.content)).active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert "시험" in headers
    col = headers.index("시험")
    body_rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    assert body_rows and all(r[col] == "시험" for r in body_rows)


def test_dhr_excel_export_has_no_test_marker():
    """배합일지 Excel 은 시험 기록도 그대로 출력하되 '시험' 표식이 **없다**(§5 · 외부 제출용)."""
    from openpyxl import load_workbook

    client, csrf = _mgmt_client()
    product = "일지확인" + _uid()
    _worker_session(client, csrf, "일지작업" + _uid())
    res = _save_test(
        client, csrf, product, rows=[("원료A", 30, 30, "TLA"), ("원료B", 70, 70, "TLB")]
    )
    assert res.status_code == 200, res.text
    record_id = res.json()["id"]

    export = client.get(f"/api/blend/records/{record_id}/export")
    assert export.status_code == 200, export.text
    wb = load_workbook(io.BytesIO(export.content))
    texts = [
        str(cell.value)
        for sheet in wb.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]
    assert not any("시험" in t for t in texts), [t for t in texts if "시험" in t]
    # 출력 자체는 정상 — 제품 LOT 과 자재명이 들어 있다.
    assert any(res.json()["product_lot"] in t for t in texts)


# ── §6 격리 ───────────────────────────────────────────────────────
def _seed_pair(client, csrf, product, worker):
    """같은 제품의 정식 1건 + 시험 1건. (정식 총량 100g, 시험 총량 100g)

    자재명은 이 호출마다 고유하게 만든다 — 자재별 집계(material_usage·불출량)는 제품과
    무관하게 이름으로 합산하므로, 이름이 겹치면 같은 DB 의 다른 테스트 값이 섞인다.
    반환: (recipe_id, 정식 기록, 시험 기록, (자재A명, 자재B명))
    """
    suffix = _uid()
    mat_a, mat_b = "원료A" + suffix, "원료B" + suffix
    rid = _import_recipe(client, csrf, product, [(mat_a, 60), (mat_b, 40)])
    prod = client.post("/api/blend/records", json={
        "recipe_id": rid, "product_name": product, "worker": "무시됨",
        "work_date": _WORK_DATE, "total_amount": 100,
        "details": [
            {"material_name": mat_a, "actual_amount": 60, "material_lot": "LA"},
            {"material_name": mat_b, "actual_amount": 40, "material_lot": "LB"},
        ],
    }, headers=csrf())
    assert prod.status_code == 200, prod.text
    test_res = _save_test(
        client, csrf, product,
        rows=[(mat_a, 30, 30, "TLA"), (mat_b, 70, 70, "TLB")],
        base_recipe_id=rid,
    )
    assert test_res.status_code == 200, test_res.text
    return rid, prod.json(), test_res.json(), (mat_a, mat_b)


def test_dashboard_excludes_test_records():
    """대시보드 집계(기간 요약·추세·최근)에서 시험은 빠진다(§6)."""
    client, csrf = _mgmt_client()
    product = "대시확인" + _uid()
    _worker_session(client, csrf, "대시작업" + _uid())
    _rid, _prod, test_rec, _mats = _seed_pair(client, csrf, product, "대시작업")

    params = {"from": _WORK_DATE, "to": _WORK_DATE}
    summary = client.get("/api/dashboard/summary", params=params)
    assert summary.status_code == 200, summary.text
    # 같은 날 다른 테스트의 기록도 같은 DB 에 있으니 절대값 대신 '정식만' 을 확인한다.
    trend = client.get("/api/dashboard/trend", params=params).json()["points"]
    day = next(p for p in trend if p["date"] == _WORK_DATE)
    excluded = client.get(
        "/api/blend/records",
        params={"start_date": _WORK_DATE, "end_date": _WORK_DATE, "test": "exclude",
                "limit": 1000},
    ).json()
    assert summary.json()["blend_count"] == excluded["total_available"]
    assert day["blend_count"] == excluded["total_available"]

    recent = client.get("/api/dashboard/recent", params={"limit": 50}).json()["items"]
    assert test_rec["product_lot"] not in [r["product_lot"] for r in recent]


def test_dashboard_export_excludes_test_records():
    """대시보드 보고서 Excel 도 화면과 같은 기준(시험 제외)."""
    from openpyxl import load_workbook

    client, csrf = _mgmt_client()
    product = "대시엑셀" + _uid()
    _worker_session(client, csrf, "대시엑셀작업" + _uid())
    _seed_pair(client, csrf, product, "대시엑셀작업")

    params = {"from": _WORK_DATE, "to": _WORK_DATE}
    res = client.get("/api/dashboard/export", params=params)
    assert res.status_code == 200, res.text
    ws = load_workbook(io.BytesIO(res.content)).active
    values = [
        [c.value for c in row] for row in ws.iter_rows()
    ]
    count_row = next(r for r in values if r and r[0] == "배합 건수")
    summary = client.get("/api/dashboard/summary", params=params).json()
    assert count_row[1] == summary["blend_count"]


def test_insight_analysis_and_material_usage_exclude_test():
    """배합 분석(/insight)의 지표·제품·자재·배치 상세에서 시험은 빠진다(§6)."""
    client, csrf = _mgmt_client()
    product = "분석확인" + _uid()
    _worker_session(client, csrf, "분석작업" + _uid())
    _rid, _prod, test_rec, (mat_a, _mat_b) = _seed_pair(client, csrf, product, "분석작업")
    params = {"start_date": _WORK_DATE, "end_date": _WORK_DATE}

    analysis = client.get("/api/blend/analysis", params=params)
    assert analysis.status_code == 200, analysis.text
    data = analysis.json()
    row = next(p for p in data["products"] if p["product_name"] == product)
    # 정식 1건만 — 시험까지 셌으면 2건·200g 이 된다.
    assert row["batch_count"] == 1
    assert row["total_amount"] == 100.0

    usage = client.get("/api/blend/material-usage", params=params).json()["items"]
    by_name = {i["material_name"]: i for i in usage}
    # 시험이 섞이면 이 자재는 60+30=90g 이 된다.
    assert by_name[mat_a]["total_actual"] == 60.0

    product_usage = client.get("/api/blend/product-usage", params=params).json()["items"]
    prow = next(p for p in product_usage if p["product_name"] == product)
    assert prow["batch_count"] == 1

    batch = client.get(
        "/api/blend/batch-details", params={**params, "product": product}
    ).json()["items"]
    assert test_rec["product_lot"] not in {b["product_lot"] for b in batch}


def test_public_material_usage_includes_test():
    """자재 불출량 공개 API 는 시험을 **포함**한다(§6 — 실제 소모라 재고에 반영)."""
    client, csrf = _mgmt_client()
    product = "불출확인" + _uid()
    _worker_session(client, csrf, "불출작업" + _uid())
    _rid, _prod, test_rec, (mat_a, _mat_b) = _seed_pair(client, csrf, product, "불출작업")
    params = {"start_date": _WORK_DATE, "end_date": _WORK_DATE}

    # 공개 API 는 내부망 IP 로만 열린다 — 사설 IP 위장 클라이언트(공개 API 테스트 패턴).
    import src.main as mainmod
    from fastapi.testclient import TestClient

    internal = TestClient(mainmod.app, client=("192.168.11.108", 50000))

    agg = internal.get("/api/public/material-usage", params=params)
    assert agg.status_code == 200, agg.text
    by_name: dict[str, float] = {}
    for item in agg.json()["items"]:
        by_name[item["material_name"]] = (
            by_name.get(item["material_name"], 0.0) + item["total_actual"]
        )
    # 정식 60 + 시험 30 = 90 — 시험이 빠지면 60 이 된다.
    assert by_name.get(mat_a, 0.0) == 90.0

    details = internal.get("/api/public/material-usage/details", params=params)
    assert details.status_code == 200, details.text
    lots = {d["product_lot"] for d in details.json()["items"]}
    assert test_rec["product_lot"] in lots


def test_lot_history_excludes_test():
    """LOT 이력(교체 타임라인)에서 시험은 빠진다(§6) — 시험 LOT 이 교체 사건이 되면 거짓이다."""
    client, csrf = _mgmt_client()
    product = "이력확인" + _uid()
    _worker_session(client, csrf, "이력작업" + _uid())
    rid = _import_recipe(client, csrf, product, [("원료A", 60), ("원료B", 40)])
    prod = client.post("/api/blend/records", json={
        "recipe_id": rid, "product_name": product, "worker": "무시됨",
        "work_date": _WORK_DATE, "total_amount": 100,
        "details": [
            {"material_name": "원료A", "actual_amount": 60, "material_lot": "LA-1"},
            {"material_name": "원료B", "actual_amount": 40, "material_lot": "LB-1"},
        ],
    }, headers=csrf())
    assert prod.status_code == 200, prod.text
    # 시험에서 원료A 를 다른 LOT 으로 계량 — 정식 이력에 '교체'로 잡히면 안 된다.
    test_res = _save_test(
        client, csrf, product,
        rows=[("원료A", 30, 30, "TEST-LOT-A"), ("원료B", 70, 70, "LB-1")],
        base_recipe_id=rid, work_date="2026-09-19",
    )
    assert test_res.status_code == 200, test_res.text

    fams = client.get("/api/blend/lot-history/families")
    assert fams.status_code == 200, fams.text
    family = next(
        (f for f in fams.json()["items"] if f["label"] == product), None
    )
    assert family is not None, fams.text
    hist = client.get("/api/blend/lot-history", params={"family": family["key"]})
    assert hist.status_code == 200, hist.text
    payload = hist.json()
    blob = str(payload)
    assert "TEST-LOT-A" not in blob
    assert test_res.json()["product_lot"] not in blob


def test_material_lot_trace_includes_test_with_flag():
    """자재 LOT 역추적은 시험을 포함하고 is_test 를 실어 준다(§6 — 리콜 추적)."""
    client, csrf = _mgmt_client()
    product = "추적확인" + _uid()
    _worker_session(client, csrf, "추적작업" + _uid())
    lot = "TRACE" + _uid()
    res = _save_test(client, csrf, product, rows=[("원료A", 10, 10, lot)])
    assert res.status_code == 200, res.text

    trace = client.get("/api/blend/material-lot-trace", params={"lot": lot})
    assert trace.status_code == 200, trace.text
    items = trace.json()["items"]
    assert items and all(int(i["is_test"]) == 1 for i in items)


# ── §7 페이지 ─────────────────────────────────────────────────────
def test_blend_test_page_renders_blend_screen():
    """GET /blend/test — /blend 와 같은 작업자 가드로 배합 화면을 시험 모드로 렌더(§7).

    TODO(패키지 B): 템플릿이 `#blend-entry-mode` 에 `data-test-mode="1"` 을 붙인 뒤
    그 속성까지 여기서 확인한다(지금은 템플릿이 test_mode 를 읽지 않아 속성이 없다).
    """
    client, csrf = _mgmt_client()
    _worker_session(client, csrf, "페이지작업" + _uid())
    res = client.get("/blend/test")
    assert res.status_code == 200, res.text
    # blend.html 이 렌더됐다는 증거(배합 화면 전용 자산).
    assert "blend.js" in res.text
    plain = client.get("/blend")
    assert plain.status_code == 200, plain.text


def test_blend_test_page_requires_worker_session():
    """작업자 세션이 없으면 /blend 와 똑같이 작업자 로그인으로 보낸다."""
    client = _client()
    res = client.get("/blend/test", follow_redirects=False)
    assert res.status_code == 303
    assert "/blend/login" in res.headers["location"]
    assert "next=/blend/test" in res.headers["location"]


def test_production_path_unchanged_smoke():
    """§4 마지막 문장 — is_test=false 경로는 한 줄도 바뀌지 않는다(대표 회귀 1건)."""
    client, csrf = _mgmt_client()
    product = "정식회귀" + _uid()
    _worker_session(client, csrf, "정식작업" + _uid())
    rid = _import_recipe(client, csrf, product, [("원료A", 60), ("원료B", 40)])
    res = client.post("/api/blend/records", json={
        "recipe_id": rid, "product_name": product, "worker": "무시됨",
        "work_date": date.today().isoformat(), "total_amount": 200,
        "details": [
            {"material_name": "원료A", "actual_amount": 120, "material_lot": "LA"},
            {"material_name": "원료B", "actual_amount": 80, "material_lot": "LB"},
        ],
    }, headers=csrf())
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["is_test"] is False
    assert body["base_recipe_id"] is None
    assert body["recipe_id"] == rid
    assert body["total_amount"] == 200.0
    assert not body["product_lot"].startswith("T-")
    assert [d["theory_amount"] for d in body["details"]] == [120.0, 80.0]
